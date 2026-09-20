"""Download linked documents and pull their text out.

Scrapy's FilesPipeline does the fetching (it already handles dedupe, retries and
the media-request plumbing); this module adds the per-job caps and the text
extraction that makes a downloaded PDF useful rather than just stored.

Output goes to `documents.jsonl` next to `pages.jsonl`, one record per file:
``{source_url, filename, page_count, markdown, content_hash}``.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
from pathlib import Path
from typing import Any, Iterator, TextIO

import scrapy
from itemadapter import ItemAdapter
from scrapy import Spider
from scrapy.pipelines.files import FilesPipeline

from app.config import get_settings

logger = logging.getLogger(__name__)

# Caps, per the feature spec. A single 200MB brochure would otherwise stall a
# crawl and fill the disk for no extraction benefit.
MAX_FILE_MB = 20
MAX_FILES_PER_JOB = 50
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024

# Extracted text is bounded too: a 500-page manual is not chatbot context.
MAX_DOCUMENT_CHARS = 400_000

_WS_RUN = re.compile(r"[ \t]{2,}")
_BLANK_RUN = re.compile(r"\n{3,}")


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    text = _WS_RUN.sub(" ", text)
    return _BLANK_RUN.sub("\n\n", text).strip()


def content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# -- extraction --------------------------------------------------------------


def extract_pdf(path: Path) -> tuple[str, int]:
    """(markdown, page_count) via PyMuPDF. Returns ("", 0) on an unreadable file."""
    try:
        import pymupdf
    except ImportError:  # pragma: no cover - dependency is declared
        logger.warning("PyMuPDF is not installed; skipping PDF text extraction")
        return "", 0

    try:
        with pymupdf.open(path) as document:
            pages = [page.get_text("text") for page in document]
            count = document.page_count
    except Exception as exc:  # noqa: BLE001 - a corrupt PDF must not fail the crawl
        logger.warning("Could not read PDF %s: %s", path.name, exc)
        return "", 0

    # A page break is a meaningful boundary in the extracted text.
    return _tidy("\n\n".join(pages))[:MAX_DOCUMENT_CHARS], count


def extract_docx(path: Path) -> tuple[str, int]:
    """(markdown, 0) via python-docx. Word has no fixed page count to report."""
    try:
        import docx
    except ImportError:  # pragma: no cover - dependency is declared
        logger.warning("python-docx is not installed; skipping DOCX text extraction")
        return "", 0

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        logger.warning("Could not read DOCX %s: %s", path.name, exc)
        return "", 0

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # Map Word's heading styles onto markdown so structure survives.
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit())
            blocks.append(f"{'#' * min(int(level or 2), 6)} {text}")
        else:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    return _tidy("\n\n".join(blocks))[:MAX_DOCUMENT_CHARS], 0


# Extension -> extractor. xlsx/pptx are downloaded and recorded, but we do not
# claim to extract them; their record carries an empty markdown body.
EXTRACTORS = {"pdf": extract_pdf, "docx": extract_docx}


def extract_document(path: Path, extension: str) -> tuple[str, int]:
    extractor = EXTRACTORS.get(extension.lower())
    if extractor is None:
        return "", 0
    return extractor(path)


# -- pipeline ----------------------------------------------------------------


class DocumentsPipeline(FilesPipeline):
    """Fetch `document_links`, extract their text, write documents.jsonl.

    Enabled per job via the `download_files` flag; when it is off the pipeline
    passes every item straight through without opening a file.
    """

    # Scrapy >= 2.12 passes `crawler` as a keyword-only argument and no longer
    # takes `settings`. FILES_STORE must already be set by then — the worker
    # passes the per-job path with `-s FILES_STORE=...` — or FilesPipeline
    # raises NotConfigured and Scrapy drops the pipeline.
    def __init__(self, store_uri: str, download_func: Any = None, **kwargs: Any):
        super().__init__(store_uri, download_func=download_func, **kwargs)
        self._enabled = False
        self._handle: TextIO | None = None
        self._requested: set[str] = set()
        self._written: set[str] = set()
        self._store_root = Path(str(store_uri))
        self.path: Path | None = None

    def open_spider(self, spider: Spider) -> None:
        # MediaPipeline.open_spider takes the spider from the crawler now and
        # warns if one is passed in.
        super().open_spider()
        self._enabled = bool(getattr(spider, "download_files", False))
        if not self._enabled:
            return
        settings = get_settings()
        job_dir = settings.job_data_dir(
            getattr(spider, "site", ""), getattr(spider, "job_id", "")
        )
        job_dir.mkdir(parents=True, exist_ok=True)
        self.path = job_dir / "documents.jsonl"
        # Append: a resumed crawl adds to the same file (see JsonLinesPipeline).
        self._written = _existing_sources(self.path)
        self._handle = self.path.open("a", encoding="utf-8")
        logger.info(
            "Downloading documents to %s (max %s files, %sMB each)",
            self._store_root,
            MAX_FILES_PER_JOB,
            MAX_FILE_MB,
        )

    def close_spider(self, spider: Spider) -> None:
        # MediaPipeline defines open_spider but no close_spider, so there is
        # deliberately no super() call here.
        if self._handle is not None:
            self._handle.close()
            self._handle = None

    def get_media_requests(self, item: Any, info: Any) -> Iterator[scrapy.Request]:
        if not self._enabled:
            return
        adapter = ItemAdapter(item)
        for entry in adapter.get("document_links") or []:
            url = entry.get("url") if isinstance(entry, dict) else str(entry)
            if not url or url in self._requested:
                continue
            if len(self._requested) >= MAX_FILES_PER_JOB:
                logger.info("Document cap of %s reached for this job", MAX_FILES_PER_JOB)
                return
            self._requested.add(url)
            yield scrapy.Request(
                url,
                # Enforced by the downloader, so an oversized file is abandoned
                # mid-transfer rather than after it has landed on disk.
                meta={"download_maxsize": MAX_FILE_BYTES, "playwright": False},
            )

    def item_completed(self, results: list[Any], item: Any, info: Any) -> Any:
        if not self._enabled or self._handle is None:
            return item

        downloaded: list[dict[str, Any]] = []
        for ok, result in results:
            if not ok or not isinstance(result, dict):
                continue
            stored = self._store_root / result["path"]
            source_url = str(result.get("url") or "")
            if source_url in self._written or not stored.exists():
                continue

            extension = stored.suffix.lstrip(".").lower()
            markdown, page_count = extract_document(stored, extension)
            record = {
                "source_url": source_url,
                "filename": stored.name,
                "page_count": page_count,
                "markdown": markdown,
                "content_hash": content_hash(markdown) if markdown else "",
                "size_bytes": stored.stat().st_size,
                "ext": extension,
            }
            self._written.add(source_url)
            self._handle.write(json.dumps(record, ensure_ascii=False) + "\n")
            self._handle.flush()
            downloaded.append(
                {"filename": record["filename"], "source_url": source_url}
            )
            logger.info(
                "Document %s: %s chars, %s pages", stored.name, len(markdown), page_count
            )

        if downloaded:
            adapter = ItemAdapter(item)
            adapter["documents_downloaded"] = len(downloaded)
        return item


def _existing_sources(path: Path) -> set[str]:
    """Source URLs already in documents.jsonl, so a resume does not redo them."""
    if not path.exists():
        return set()
    found: set[str] = set()
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                found.add(str(json.loads(line).get("source_url") or ""))
            except json.JSONDecodeError:
                continue
    found.discard("")
    return found
